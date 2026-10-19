#!/usr/bin/env python3
"""make_docx.py -- article.md -> article_flux_oral_ecrit.docx in the style of Studii de lingvistica
(pandoc with the journal's template as reference document, then python-docx post-processing:
Times New Roman, title 14 bold centred, authors right, abstracts/keywords/examples/references 11,
text 12 justified, captions centred, margins 2.5 cm, bordered tables)."""
import re
import subprocess
import docx
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC, REF, OUT = "article.md", "feuille-de-style-SL.docx", "article_flux_oral_ecrit.docx"
subprocess.run(["pandoc", SRC, "--reference-doc", REF, "-o", OUT], check=True)
d = docx.Document(OUT)
TNR = "Times New Roman"


def font(p, size, bold=None):
    for r in p.runs:
        r.font.name = TNR
        r._element.rPr.rFonts.set(qn("w:eastAsia"), TNR)
        r.font.size = Pt(size)
        if bold is not None:
            r.font.bold = bold


for st in ("Normal", "Body Text", "Heading 1", "Title"):
    try:
        d.styles[st].font.name = TNR
    except KeyError:
        pass
for s in d.sections:
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)

state = "start"
for p in d.paragraphs:
    t = p.text.strip()
    sty = p.style.name
    if sty == "Title":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p, 14, True)
        continue
    if state == "start" and t.startswith("Sylvain Kahane"):
        p.style = d.styles["Body Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        font(p, 12)
        state = "abstract"
        continue
    if re.match(r"^(Résumé|Mots-clés|Abstract|Key words|Rezumat|Cuvinte-cheie)\s?:", t):
        p.style = d.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        font(p, 11)
        p.paragraph_format.space_after = Pt(6)
        continue
    if re.match(r"^\d+\.\d+\.\s", t) or (sty == "Heading 1"):
        p.style = d.styles["Heading 1"]
        font(p, 12, True)
        state = "body"
        continue
    if t.startswith("Références bibliographiques"):
        p.style = d.styles["Heading 1"]
        font(p, 12, True)
        state = "refs"
        continue
    if re.match(r"^\(\d\)", t):
        p.style = d.styles["List Paragraph"] if "List Paragraph" in [s.name for s in d.styles] else d.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        font(p, 11)
        continue
    if re.match(r"^(Figure|Tableau) \d", t):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p, 11)
        continue
    if state == "refs" and t:
        p.style = d.styles["Normal"]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        font(p, 11)
        continue
    if p._element.xpath(".//w:drawing"):
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    if t and state in ("body", "abstract"):
        p.style = d.styles["Body Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        font(p, 12)


def borders(tbl):
    tblPr = tbl._element.tblPr
    b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{e}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        b.append(el)
    tblPr.append(b)


for tbl in d.tables:
    borders(tbl)
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                font(p, 10)
# footnotes
try:
    fn = d.part.package.part_related_by  # noqa
except Exception:
    pass
for rel in d.part.rels.values():
    if "footnotes" in rel.reltype:
        root = rel.target_part.element if hasattr(rel.target_part, "element") else None
        if root is None:
            from lxml import etree
            root = etree.fromstring(rel.target_part.blob)
        for r in root.iter(qn("w:r")):
            rPr = r.find(qn("w:rPr"))
            if rPr is None:
                rPr = OxmlElement("w:rPr")
                r.insert(0, rPr)
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), TNR)
            fonts.set(qn("w:hAnsi"), TNR)
            rPr.append(fonts)
            sz = OxmlElement("w:sz")
            sz.set(qn("w:val"), "20")
            rPr.append(sz)
        if hasattr(rel.target_part, "_blob"):
            from lxml import etree
            rel.target_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
d.save(OUT)
print("written", OUT)
